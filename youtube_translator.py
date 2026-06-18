#!/usr/bin/env python3
"""YouTube動画の字幕を取得し、日本語に翻訳・5回校正してWord/TXTで保存するアプリ"""

import os
import re
import sys
import argparse
from pathlib import Path

from youtube_transcript_api import YouTubeTranscriptApi
from youtube_transcript_api.formatters import TextFormatter
from docx import Document
import anthropic


def extract_video_id(url: str) -> str:
    patterns = [
        r'(?:v=|/v/|youtu\.be/)([a-zA-Z0-9_-]{11})',
        r'^([a-zA-Z0-9_-]{11})$',
    ]
    for p in patterns:
        m = re.search(p, url)
        if m:
            return m.group(1)
    raise ValueError(f"無効なYouTube URL: {url}")


def get_transcript(video_id: str) -> str:
    try:
        transcript = YouTubeTranscriptApi.get_transcript(video_id, languages=['en', 'ja', 'ko', 'zh', 'es', 'fr', 'de'])
    except Exception:
        transcript = YouTubeTranscriptApi.get_transcript(video_id)
    formatter = TextFormatter()
    return formatter.format_transcript(transcript)


def get_video_title(video_id: str) -> str:
    try:
        import urllib.request
        import json
        url = f"https://www.youtube.com/oembed?url=https://www.youtube.com/watch?v={video_id}&format=json"
        with urllib.request.urlopen(url, timeout=10) as resp:
            data = json.loads(resp.read().decode())
            return data.get("title", video_id)
    except Exception:
        return video_id


def translate_to_japanese(client: anthropic.Anthropic, text: str) -> str:
    resp = client.messages.create(
        model="claude-sonnet-4-6",
        max_tokens=8192,
        messages=[{
            "role": "user",
            "content": f"以下のテキストを自然な日本語に翻訳してください。翻訳のみを出力し、説明や注釈は不要です。\n\n{text}"
        }]
    )
    return resp.content[0].text


def proofread(client: anthropic.Anthropic, text: str, iteration: int) -> str:
    resp = client.messages.create(
        model="claude-sonnet-4-6",
        max_tokens=8192,
        messages=[{
            "role": "user",
            "content": (
                f"以下の日本語翻訳文を校正してください（校正 {iteration}/5回目）。\n"
                "・誤訳や不自然な表現を修正\n"
                "・文法・句読点の誤りを修正\n"
                "・読みやすさと正確性を向上\n"
                "校正後の文章のみを出力してください。\n\n"
                f"{text}"
            )
        }]
    )
    return resp.content[0].text


def save_as_txt(text: str, filepath: Path):
    filepath.write_text(text, encoding="utf-8")


def save_as_docx(text: str, filepath: Path, title: str):
    doc = Document()
    doc.add_heading(title, level=1)
    for para in text.split("\n"):
        doc.add_paragraph(para)
    doc.save(str(filepath))


def sanitize_filename(name: str) -> str:
    return re.sub(r'[\\/*?:"<>|]', "_", name)[:100]


def process_video(client: anthropic.Anthropic, url: str, output_dir: Path, index: int):
    video_id = extract_video_id(url)
    title = get_video_title(video_id)
    safe_title = sanitize_filename(title)

    print(f"\n{'='*60}")
    print(f"[{index}] 処理中: {title}")
    print(f"    URL: {url}")
    print(f"{'='*60}")

    print("  字幕を取得中...")
    transcript = get_transcript(video_id)
    print(f"  字幕取得完了 ({len(transcript)} 文字)")

    print("  日本語に翻訳中...")
    translated = translate_to_japanese(client, transcript)
    print("  翻訳完了")

    for i in range(1, 6):
        print(f"  校正中 ({i}/5)...")
        translated = proofread(client, translated, i)
    print("  校正完了")

    txt_path = output_dir / f"{safe_title}.txt"
    docx_path = output_dir / f"{safe_title}.docx"

    save_as_txt(translated, txt_path)
    save_as_docx(translated, docx_path, title)

    print(f"  保存完了:")
    print(f"    TXT:  {txt_path}")
    print(f"    DOCX: {docx_path}")


def main():
    parser = argparse.ArgumentParser(description="YouTube動画の字幕を日本語に翻訳・校正して保存")
    parser.add_argument("--urls", nargs="+", help="YouTube URL（最大5件）")
    parser.add_argument("--output", "-o", default="./output", help="出力先フォルダ（デフォルト: ./output）")
    args = parser.parse_args()

    if args.urls:
        urls = args.urls
    else:
        print("YouTube URLを入力してください（最大5件、空行で入力終了）:")
        urls = []
        for i in range(5):
            url = input(f"  URL {i+1} (空行で終了): ").strip()
            if not url:
                break
            urls.append(url)

    if not urls:
        print("URLが入力されていません。終了します。")
        sys.exit(1)

    if len(urls) > 5:
        print("エラー: URLは最大5件までです。")
        sys.exit(1)

    api_key = os.environ.get("ANTHROPIC_API_KEY")
    if not api_key:
        print("エラー: 環境変数 ANTHROPIC_API_KEY を設定してください。")
        sys.exit(1)

    output_dir = Path(args.output)
    output_dir.mkdir(parents=True, exist_ok=True)

    client = anthropic.Anthropic(api_key=api_key)

    print(f"\n処理対象: {len(urls)} 件")
    print(f"出力先: {output_dir.resolve()}")

    for i, url in enumerate(urls, 1):
        try:
            process_video(client, url, output_dir, i)
        except Exception as e:
            print(f"\n  エラー [{i}]: {e}")

    print(f"\n{'='*60}")
    print("全処理完了！")
    print(f"{'='*60}")


if __name__ == "__main__":
    main()
